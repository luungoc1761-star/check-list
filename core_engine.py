"""
core_engine.py
Xử lý logic đối soát BOM (Production BOM List) và Tồn kho (Stock 61 / Stock Balance).
- Nhận diện Product Code:
    + Bắt đầu bằng 1: 4 số đầu là mã đơn hàng, 2 số tiếp theo (19: Trái, 20: Phải).
    + Bắt đầu bằng 8: bỏ số 8, 4 số tiếp theo là mã đơn hàng, số thứ 6 (1: Trái, 2: Phải).
- Cột Component Code đối soát với Item code / Stock Code trong Stock 61 (hoặc Stock Balance).
- Quản lý trừ kho theo cơ chế FIFO theo thứ tự đơn hàng nhập vào hàng chờ.
- Xuất file Excel định dạng chuyên nghiệp với tô màu xanh dòng OK, đỏ dòng THIẾU.
"""

import io
import re
import uuid
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def clean_order_title(desc, order_no=""):
    """
    Loại bỏ các từ chỉ vế Trái / Phải (LHF, RHF, LEFT, RIGHT, Side Assy, Side...)
    để hiển thị tên sản phẩm đại diện nguyên đơn gọn gàng, không bị dính chữ Trái/Phải.
    """
    if not desc or str(desc).strip().lower() in ('', 'nan', 'none'):
        return f"Sản phẩm {order_no}" if order_no else "Sản phẩm"
    
    cleaned = str(desc).strip()
    # Xóa các mã tiền tố kiểu "8167110301 - " hoặc "8184610004 -"
    cleaned = re.sub(r'^\d{8,12}\s*[-_:]\s*', '', cleaned)
    # Xóa các từ LHF, RHF, LH, RH, LEFT, RIGHT, Side Assy, Side
    cleaned = re.sub(r'\b(LHF|RHF|LH|RH|LEFT|RIGHT|Side\s*Assy|Side)\b', '', cleaned, flags=re.IGNORECASE)
    # Dọn dẹp khoảng trắng thừa và dấu gạch nối lẻ loi
    cleaned = re.sub(r'[-_]\s*$', '', cleaned)
    cleaned = re.sub(r'^\s*[-_]', '', cleaned)
    cleaned = re.sub(r'\s+', ' ', cleaned).strip()
    
    if not cleaned or cleaned.isdigit():
        return f"Sản phẩm {order_no}"
    return cleaned


def parse_product_code(code):
    """
    Phân tích mã sản phẩm / Model:
    - Bắt đầu bằng 1: 4 số đầu là đơn hàng, 2 số tiếp theo là vế (19: Trái, 20: Phải)
    - Bắt đầu bằng 8: bỏ 8, 4 số tiếp theo là đơn hàng, số tiếp theo là vế (1: Trái, 2: Phải)
    """
    code_str = str(code).strip()
    if code_str.startswith('1'):
        order = code_str[:4]
        side_raw = code_str[4:6]
        side = 'Trái' if side_raw == '19' else ('Phải' if side_raw == '20' else f'Vế {side_raw}')
        variant = code_str[6:]
        return order, side, variant
    elif code_str.startswith('8'):
        order = code_str[1:5]
        side_raw = code_str[5] if len(code_str) > 5 else ''
        side = 'Trái' if side_raw == '1' else ('Phải' if side_raw == '2' else f'Vế {side_raw}')
        variant = code_str[6:] if len(code_str) > 6 else ''
        return order, side, variant
    else:
        return code_str[:4], 'Khác', code_str[4:]


def find_plan_header_row_from_df(df_raw, max_rows=20):
    """
    Tìm chỉ số dòng chứa tiêu đề file kế hoạch (MODEL và Qty).
    """
    model_kws = ['model', 'product code', 'product_code', 'mã sp', 'mã hàng', 'mã model', 'mã']
    qty_kws = ['qty', 'quantity', 'số lượng', 'sl']
    for i in range(min(max_rows, len(df_raw))):
        row_str = ' '.join([str(x).lower().strip() for x in df_raw.iloc[i].values if pd.notna(x)])
        if any(k in row_str for k in model_kws) and any(k in row_str for k in qty_kws):
            return i
    return 0


def load_plan_data(file_source):
    """
    Đọc file Kế hoạch sản xuất (.xlsx):
    - Lấy dữ liệu 2 cột MODEL và Qty (và Due date nếu có).
    - Giải mã cột MODEL ra Mã đơn hàng (Order_No), Vế (Trái/Phải) và Phiên bản (Variant).
    """
    df_raw = pd.read_excel(file_source, header=None, nrows=20)
    h_idx = find_plan_header_row_from_df(df_raw)
    if hasattr(file_source, 'seek'):
        file_source.seek(0)
    df = pd.read_excel(file_source, skiprows=h_idx)
    
    col_map = {}
    for c in df.columns:
        c_clean = str(c).strip().lower()
        if any(k in c_clean for k in ['model', 'product code', 'product_code', 'mã sp', 'mã hàng', 'mã model']):
            if 'MODEL' not in col_map.values():
                col_map[c] = 'MODEL'
        elif any(k in c_clean for k in ['qty', 'quantity', 'số lượng', 'sl']):
            if 'Qty' not in col_map.values():
                col_map[c] = 'Qty'
        elif any(k in c_clean for k in ['due date', 'due_date', 'ngày', 'hạn', 'date']):
            if 'Due_Date' not in col_map.values():
                col_map[c] = 'Due_Date'
                
    df = df.rename(columns=col_map)
    if 'MODEL' not in df.columns or 'Qty' not in df.columns:
        raise ValueError('Không tìm thấy cột MODEL hoặc Qty trong file Kế hoạch.')
        
    df = df[pd.notna(df['MODEL']) & pd.notna(df['Qty'])].copy()
    df['MODEL'] = df['MODEL'].astype(str).str.strip()
    df['Qty'] = pd.to_numeric(df['Qty'], errors='coerce').fillna(0)
    df = df[df['Qty'] > 0].copy()
    
    parsed = [parse_product_code(m) for m in df['MODEL']]
    df['Order_No'] = [p[0] for p in parsed]
    df['Side'] = [p[1] for p in parsed]
    df['Variant'] = [p[2] for p in parsed]
    
    if 'Due_Date' in df.columns:
        df['Due_Date'] = df['Due_Date'].astype(str).str.slice(0, 10)
    else:
        df['Due_Date'] = ''
        
    return df


def extract_orders_from_plan(df_plan):
    """
    Chuyển đổi dữ liệu từ File Kế hoạch thành danh sách hàng chờ đối soát FIFO.
    Giữ nguyên thứ tự ưu tiên từ trên xuống dưới của File Kế hoạch.
    Mỗi đơn hàng tự động ghép vế Trái và Phải từ các dòng Model tương ứng.
    """
    orders_queue = []
    seen = {}
    
    for idx, row in df_plan.iterrows():
        ord_no = str(row['Order_No']).strip()
        side = str(row['Side']).strip()
        qty = float(row['Qty'])
        variant = str(row['Variant']).strip() if pd.notna(row['Variant']) else ''
        model = str(row['MODEL']).strip()
        due_date = str(row.get('Due_Date', '')).strip()
        
        if ord_no not in seen:
            item = {
                'id': f'plan_{ord_no}_{len(orders_queue)+1}_{uuid.uuid4().hex[:6]}',
                'order_no': ord_no,
                'qty': qty,
                'qty_left': qty if side == 'Trái' else 0.0,
                'qty_right': qty if side == 'Phải' else 0.0,
                'variant': variant,
                'due_date': due_date,
                'models': {side: model},
                'source': 'Kế hoạch',
                'created_at': due_date if due_date else 'Kế hoạch'
            }
            seen[ord_no] = item
            orders_queue.append(item)
        else:
            item = seen[ord_no]
            if side == 'Trái':
                item['qty_left'] = qty
            elif side == 'Phải':
                item['qty_right'] = qty
            item['models'][side] = model
            item['qty'] = max(item['qty_left'], item['qty_right'])
            if variant and not item.get('variant'):
                item['variant'] = variant
            if due_date and not item.get('due_date'):
                item['due_date'] = due_date
                
    return orders_queue


def find_header_row_from_df(df_raw, keywords, max_rows=20):
    """Tìm chỉ số dòng chứa các từ khóa tiêu đề."""
    for i in range(min(max_rows, len(df_raw))):
        row_str = " ".join([str(x).lower().strip() for x in df_raw.iloc[i].values if pd.notna(x)])
        if all(kw.lower() in row_str for kw in keywords):
            return i
    return 0


def load_bom_data(file_source):
    """
    Đọc dữ liệu Production BOM từ đường dẫn file hoặc file-like object (uploaded file).
    Tự động nhận diện dòng tiêu đề và chuẩn hóa cột.
    """
    df_raw = pd.read_excel(file_source, header=None, nrows=20)
    h_idx = find_header_row_from_df(df_raw, ['product code', 'component code'])
    
    if hasattr(file_source, 'seek'):
        file_source.seek(0)
        
    df = pd.read_excel(file_source, skiprows=h_idx)
    
    col_map = {}
    for c in df.columns:
        c_clean = str(c).strip().lower()
        if 'product code' in c_clean:
            col_map[c] = 'Product_Code'
        elif 'production desc' in c_clean or 'product desc' in c_clean:
            if 'Product_Desc' not in col_map.values():
                col_map[c] = 'Product_Desc'
        elif 'component code' in c_clean:
            col_map[c] = 'Component_Code'
        elif 'component desc' in c_clean:
            if 'Component_Desc' not in col_map.values():
                col_map[c] = 'Component_Desc'
        elif c_clean == 'qty' or 'qty' in c_clean:
            if 'Qty' not in col_map.values():
                col_map[c] = 'Qty'
                
    df = df.rename(columns=col_map)
    df = df[pd.notna(df['Product_Code']) & pd.notna(df['Component_Code'])].copy()
    
    df['Product_Code'] = df['Product_Code'].astype(str).str.strip()
    df['Component_Code'] = df['Component_Code'].astype(str).str.strip()
    df['Qty'] = pd.to_numeric(df['Qty'], errors='coerce').fillna(0)
    
    if 'Product_Desc' not in df.columns:
        df['Product_Desc'] = ''
    else:
        df['Product_Desc'] = df['Product_Desc'].fillna('').astype(str).str.strip()
        
    if 'Component_Desc' not in df.columns:
        df['Component_Desc'] = ''
    else:
        df['Component_Desc'] = df['Component_Desc'].fillna('').astype(str).str.strip()
        
    parsed = [parse_product_code(c) for c in df['Product_Code']]
    df['Order_No'] = [p[0] for p in parsed]
    df['Side'] = [p[1] for p in parsed]
    df['Variant'] = [p[2] for p in parsed]
    
    return df


def find_stock_header_row_from_df(df_raw, max_rows=20):
    """
    Tìm chỉ số dòng chứa tiêu đề bảng tồn kho.
    Hỗ trợ linh hoạt cả:
    - File Stock 61: 'Item code', 'Wh', 'Qty'
    - File Stock Balance: 'Stock Code', 'Warehouse', 'Qty',...
    - Các file tiếng Việt hoặc định dạng kho thông dụng khác.
    """
    code_kws = ['stock code', 'item code', 'item_code', 'component code', 'stock_code', 'mã vt', 'mã linh kiện', 'mã hàng', 'mã', 'item', 'code']
    qty_kws = ['qty', 'quantity', 'số lượng', 'sl', 'tồn kho', 'tồn']
    for i in range(min(max_rows, len(df_raw))):
        row_str = " ".join([str(x).lower().strip() for x in df_raw.iloc[i].values if pd.notna(x)])
        if any(k in row_str for k in code_kws) and any(k in row_str for k in qty_kws):
            return i
    return 0


def load_stock_data(file_source):
    """
    Đọc dữ liệu Tồn kho từ đường dẫn file hoặc file-like object (hỗ trợ Stock 61, Stock Balance,...).
    Tự động nhận diện dòng tiêu đề và chuẩn hóa cột.
    """
    df_raw = pd.read_excel(file_source, header=None, nrows=20)
    h_idx = find_stock_header_row_from_df(df_raw)
    
    if hasattr(file_source, 'seek'):
        file_source.seek(0)
        
    df = pd.read_excel(file_source, skiprows=h_idx)
    
    col_map = {}
    for c in df.columns:
        c_clean = str(c).strip().lower()
        if any(k in c_clean for k in ['stock code', 'item code', 'item_code', 'stock_code', 'component code', 'mã vt', 'mã linh kiện', 'mã hàng']) or c_clean in ('item', 'code', 'mã'):
            if 'Stock_Code' not in col_map.values():
                col_map[c] = 'Stock_Code'
        elif any(k in c_clean for k in ['warehouse', 'wh', 'kho']):
            if 'Warehouse' not in col_map.values():
                col_map[c] = 'Warehouse'
        elif 'batch' in c_clean or 'lô' in c_clean:
            if 'BATCH' not in col_map.values():
                col_map[c] = 'BATCH'
        elif 'bin' in c_clean or 'vị trí' in c_clean:
            if 'BIN' not in col_map.values():
                col_map[c] = 'BIN'
        elif c_clean == 'qty' or any(k in c_clean for k in ['qty', 'quantity', 'số lượng', 'sl', 'tồn kho', 'tồn']):
            if 'Qty' not in col_map.values():
                col_map[c] = 'Qty'
        elif 'date' in c_clean or 'ngày' in c_clean:
            if 'Date' not in col_map.values():
                col_map[c] = 'Date'
            
    df = df.rename(columns=col_map)
    if 'Stock_Code' not in df.columns:
        raise ValueError("Không tìm thấy cột Mã linh kiện / Item Code / Stock Code trong file tồn kho.")
    if 'Qty' not in df.columns:
        raise ValueError("Không tìm thấy cột Số lượng (Qty) trong file tồn kho.")
        
    df = df[pd.notna(df['Stock_Code'])].copy()
    df['Stock_Code'] = df['Stock_Code'].astype(str).str.strip()
    df = df[~df['Stock_Code'].str.lower().isin(['', 'nan', 'none'])].copy()
    
    df['Qty'] = pd.to_numeric(df['Qty'], errors='coerce').fillna(0)
    
    for c in ['BATCH', 'BIN', 'Warehouse']:
        if c in df.columns:
            df[c] = df[c].fillna('').astype(str).str.strip()
        else:
            df[c] = ''
            
    return df


def get_available_orders_info(df_bom):
    """
    Trích xuất danh sách tất cả các đơn hàng có trong BOM cùng thông tin tóm tắt:
    tên sản phẩm, số biến thể, số linh kiện, có đủ trái phải không.
    """
    orders_info = {}
    grouped = df_bom.groupby('Order_No')
    
    for order_no, group in grouped:
        prods = group[['Product_Code', 'Product_Desc', 'Side', 'Variant']].drop_duplicates()
        sides = sorted(group['Side'].unique())
        has_both_sides = ('Trái' in sides) and ('Phải' in sides)
        
        # Lấy mô tả sản phẩm đại diện (làm sạch không chứa Trái / Phải / LHF / RHF)
        desc_list = [d for d in group['Product_Desc'].unique() if d and str(d).strip().lower() not in ('', 'nan')]
        raw_desc = desc_list[0] if desc_list else f"Sản phẩm đơn hàng {order_no}"
        clean_desc = clean_order_title(raw_desc, order_no)
        
        variants = sorted(group['Variant'].unique())
        unique_components = len(group['Component_Code'].unique())
        total_bom_rows = len(group)
        
        orders_info[order_no] = {
            'order_no': order_no,
            'description': clean_desc,
            'raw_description': raw_desc,
            'sides': sides,
            'has_both_sides': has_both_sides,
            'variants': variants,
            'unique_components': unique_components,
            'total_bom_rows': total_bom_rows,
            'products': prods.to_dict(orient='records')
        }
        
    return orders_info


def calculate_inventory_allocation(orders_queue, df_bom, df_stock):
    """
    Tính toán đối soát BOM và tồn kho theo cơ chế FIFO cho danh sách đơn hàng chờ.
    Mỗi đơn hàng được tính toán:
    - Lấy đầy đủ định mức Trái và Phải từ BOM.
    - Nhân số lượng đơn hàng với định mức BOM.
    - So sánh với tồn kho khả dụng hiện tại.
    - Ưu tiên cấp liệu cho đơn hàng nhập trước, trừ dần tồn kho.
    """
    # Khởi tạo quỹ tồn kho độc lập để trừ dần
    stock_qty_pool = df_stock.groupby('Stock_Code')['Qty'].sum().to_dict()
    results = []
    
    for idx, item in enumerate(orders_queue, 1):
        order_no = str(item['order_no']).strip()
        order_qty = float(item['qty'])
        qty_left = float(item.get('qty_left', order_qty))
        qty_right = float(item.get('qty_right', order_qty))
        variant = item.get('variant', None)
        item_id = item.get('id', f"ord_{idx}")
        created_at = item.get('created_at', '')
        models_map = item.get('models', {})
        due_date = item.get('due_date', '')
        
        # Lọc dữ liệu BOM cho đơn hàng này
        order_boms = df_bom[df_bom['Order_No'] == order_no]
        
        if len(order_boms) == 0:
            results.append({
                'id': item_id,
                'stt': idx,
                'order_no': order_no,
                'order_qty': order_qty,
                'qty_left': qty_left,
                'qty_right': qty_right,
                'models': models_map,
                'due_date': due_date,
                'created_at': created_at,
                'error': f"Không tìm thấy dữ liệu BOM cho đơn hàng {order_no}",
                'total_items': 0,
                'ok_count': 0,
                'missing_count': 0,
                'summary_table': pd.DataFrame(),
                'detail_table': pd.DataFrame()
            })
            continue
            
        # Nếu đơn hàng có nhiều phiên bản (variant) mà không chỉ định rõ:
        # Chọn phiên bản đầu tiên (ví dụ '0004') để không bị cộng gộp nhân đôi định mức
        avail_variants = sorted(order_boms['Variant'].unique())
        chosen_variant = variant if (variant and variant in avail_variants) else avail_variants[0]
        
        bom_order = order_boms[order_boms['Variant'] == chosen_variant].copy()
        if len(bom_order) == 0:
            bom_order = order_boms.copy()
            
        # Lấy tên sản phẩm đại diện (làm sạch không chứa Trái / Phải)
        desc_list = [d for d in bom_order['Product_Desc'].unique() if d and str(d).strip().lower() not in ('', 'nan')]
        raw_desc = desc_list[0] if desc_list else f"Sản phẩm đơn hàng {order_no}"
        product_title = clean_order_title(raw_desc, order_no)
        
        # Gom nhóm theo từng linh kiện (Component Code)
        comp_groups = bom_order.groupby('Component_Code')
        summary_rows = []
        detail_rows = []

        def get_side_priority(side):
            s = str(side).strip().lower()
            if any(k in s for k in ['trái', 'trai', 'left', 'lh']):
                return 0
            elif any(k in s for k in ['phải', 'phai', 'right', 'rh']):
                return 1
            return 2
        
        for comp_code, cgroup in comp_groups:
            comp_desc = cgroup['Component_Desc'].iloc[0]
            
            # Định mức bên Trái (1 vế)
            left_rows = cgroup[cgroup['Side'] == 'Trái']
            left_bom_qty = float(left_rows['Qty'].sum()) if len(left_rows) > 0 else 0.0
            
            # Định mức bên Phải (1 vế)
            right_rows = cgroup[cgroup['Side'] == 'Phải']
            right_bom_qty = float(right_rows['Qty'].sum()) if len(right_rows) > 0 else 0.0
            
            # Nếu có dòng khác ngoài Trái/Phải
            other_rows = cgroup[~cgroup['Side'].isin(['Trái', 'Phải'])]
            other_bom_qty = float(other_rows['Qty'].sum()) if len(other_rows) > 0 else 0.0
            
            # Tổng định mức cho 1 bộ nguyên đơn (Trái + Phải)
            total_bom_unit = left_bom_qty + right_bom_qty + other_bom_qty
            
            # Nhân định mức Trái và Phải với số lượng đơn hàng cần chạy
            req_left = left_bom_qty * qty_left
            req_right = right_bom_qty * qty_right
            req_other = other_bom_qty * max(qty_left, qty_right)
            
            # Tổng số lượng BOM cần chạy cho nguyên đơn (Trái + Phải)
            total_required = req_left + req_right + req_other
            
            # Tồn kho khả dụng hiện có trước khi trừ cho linh kiện này
            initial_stock = float(stock_qty_pool.get(comp_code, 0.0))
            running_stock = initial_stock
            
            # Sắp xếp các dòng BOM của linh kiện theo thứ tự trừ dần: Trái (0) -> Phải (1) -> Khác (2)
            cgroup_sorted = cgroup.copy()
            cgroup_sorted['_prio'] = cgroup_sorted['Side'].apply(get_side_priority)
            cgroup_sorted = cgroup_sorted.sort_values(by=['_prio', 'Product_Code']).drop(columns=['_prio'])
            
            comp_allocated = 0.0
            comp_shortage = 0.0
            
            for _, brow in cgroup_sorted.iterrows():
                b_side = brow['Side']
                b_pcode = brow['Product_Code']
                b_pdesc = brow['Component_Desc']
                b_qty = float(brow['Qty'])
                
                # Số lượng đơn theo từng vế
                if b_side == 'Trái':
                    side_order_qty = qty_left
                elif b_side == 'Phải':
                    side_order_qty = qty_right
                else:
                    side_order_qty = max(qty_left, qty_right)
                    
                row_required = b_qty * side_order_qty  # Phần BOM hàng cần của dòng này
                
                # Tồn kho khả dụng của dòng này trước khi trừ
                row_avail_stock = max(0.0, running_stock)
                
                # Công thức tính trừ dần:
                # Trừ vế trái trước, số còn lại trừ tiếp cho vế phải, chỉ báo thiếu số lượng cho mã thiếu cuối cùng
                if running_stock >= row_required:
                    row_allocated = row_required
                    row_shortage = 0.0
                    running_stock -= row_required
                    row_status = "OK"
                    row_note = "Đủ hàng"
                else:
                    row_allocated = max(0.0, running_stock)
                    row_shortage = row_required - row_allocated
                    running_stock = 0.0
                    row_status = "THIẾU"
                    row_note = f"Thiếu {int(row_shortage) if row_shortage.is_integer() else row_shortage:g}"
                    
                comp_allocated += row_allocated
                comp_shortage += row_shortage
                
                detail_rows.append({
                    'Vế': b_side,
                    'Mã sản phẩm (Product Code)': b_pcode,
                    'ITEM CODE': comp_code,
                    'DESCRIPTION (tên mô tả)': b_pdesc,
                    'Định mức BOM': b_qty,
                    'Số lượng đơn': side_order_qty,
                    'Số lượng cần': row_required,
                    'Tồn kho khả dụng': row_avail_stock,
                    'Số lượng cấp': row_allocated,
                    'Số lượng thiếu': row_shortage,
                    'Trạng thái': row_status,
                    'Ghi chú': row_note
                })
            
            # Cập nhật số lượng tồn kho còn lại vào quỹ kho FIFO
            stock_qty_pool[comp_code] = running_stock
            
            # Bảng tổng hợp linh kiện
            if comp_shortage == 0.0:
                summary_status = "OK"
                summary_note = "Đủ hàng"
            else:
                summary_status = "THIẾU"
                summary_note = f"Thiếu {int(comp_shortage) if comp_shortage.is_integer() else comp_shortage:g}"
                
            summary_rows.append({
                'ITEM CODE': comp_code,
                'DESCRIPTION (tên mô tả)': comp_desc,
                'Định mức Trái': left_bom_qty,
                'Định mức Phải': right_bom_qty,
                'Tổng định mức (1 bộ)': total_bom_unit,
                'SL Cần Trái': req_left,
                'SL Cần Phải': req_right,
                'SL BOM cần chạy': total_required,
                'Số lượng tồn kho': initial_stock,
                'Số lượng cấp': comp_allocated,
                'Số lượng thiếu': comp_shortage,
                'Trạng thái': summary_status,
                'Ghi chú': summary_note
            })
            
        df_summary = pd.DataFrame(summary_rows)
        # Sắp xếp ưu tiên hiển thị các mã THIẾU lên đầu để dễ nhận biết
        if len(df_summary) > 0:
            df_summary = df_summary.sort_values(
                by=['Trạng thái', 'Số lượng thiếu', 'ITEM CODE'],
                ascending=[False, False, True]
            ).reset_index(drop=True)
            
        # Sắp xếp bảng chi tiết BOM: Vế Trái trước, Vế Phải sau, ưu tiên dòng THIẾU lên đầu
        df_detail = pd.DataFrame(detail_rows)
        if len(df_detail) > 0:
            df_detail['_side_prio'] = df_detail['Vế'].apply(get_side_priority)
            df_detail = df_detail.sort_values(
                by=['_side_prio', 'Trạng thái', 'ITEM CODE'],
                ascending=[True, False, True]
            ).drop(columns=['_side_prio']).reset_index(drop=True)
            
        total_items = len(df_summary)
        ok_count = int((df_summary['Trạng thái'] == 'OK').sum()) if total_items > 0 else 0
        missing_count = int((df_summary['Trạng thái'] == 'THIẾU').sum()) if total_items > 0 else 0
        
        results.append({
            'id': item_id,
            'stt': idx,
            'order_no': order_no,
            'order_qty': order_qty,
            'qty_left': qty_left,
            'qty_right': qty_right,
            'models': models_map,
            'due_date': due_date,
            'product_title': product_title,
            'variant': chosen_variant,
            'created_at': created_at,
            'total_items': total_items,
            'ok_count': ok_count,
            'missing_count': missing_count,
            'is_fully_ok': (missing_count == 0 and total_items > 0),
            'summary_table': df_summary,
            'detail_table': df_detail
        })
        
    return results


def export_order_to_excel(order_result):
    """
    Xuất kết quả đối soát của một đơn hàng ra file Excel .xlsx với định dạng chuyên nghiệp:
    - Sheet 1: Tong_Hop_Vat_Tu (Tổng hợp theo mã linh kiện)
    - Sheet 2: Chi_Tiet_BOM_Trai_Phai (Chi tiết theo từng dòng BOM Trái/Phải)
    - Dòng OK: Tô màu xanh lá nhạt (#D4EDDA), chữ xanh đậm (#155724)
    - Dòng THIẾU: Tô màu đỏ nhạt (#F8D7DA), chữ đỏ đậm (#721C24)
    - Tiêu đề header màu xanh dương sang trọng, border thanh mảnh, auto fit độ rộng cột.
    """
    wb = openpyxl.Workbook()
    
    # Sheet 1: Tổng hợp
    ws1 = wb.active
    ws1.title = "Tong_Hop_Vat_Tu"
    
    # Sheet 2: Chi tiết
    ws2 = wb.create_sheet(title="Chi_Tiet_BOM_Trai_Phai")
    
    # Định nghĩa Styles
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    
    font_ok = Font(name="Segoe UI", size=10, color="155724")
    fill_ok = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    
    font_missing = Font(name="Segoe UI", size=10, bold=True, color="721C24")
    fill_missing = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    
    font_title = Font(name="Segoe UI", size=13, bold=True, color="1E3A8A")
    font_sub = Font(name="Segoe UI", size=10, italic=True, color="4B5563")
    
    border_thin = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    
    # ===== SHEET 1 =====
    df1 = order_result['summary_table']
    order_no = order_result['order_no']
    order_qty = order_result['order_qty']
    p_title = order_result.get('product_title', '')
    
    ws1.cell(row=1, column=1, value=f"KẾT QUẢ ĐỐI SOÁT TỒN KHO - ĐƠN HÀNG: {order_no} ({p_title})").font = font_title
    ws1.cell(row=2, column=1, value=f"Số lượng sản xuất: {order_qty:g} bộ (Trái + Phải) | Tổng linh kiện: {order_result['total_items']} | ĐỦ: {order_result['ok_count']} | THIẾU: {order_result['missing_count']}").font = font_sub
    
    start_r1 = 4
    headers1 = list(df1.columns)
    for col_idx, col_name in enumerate(headers1, 1):
        cell = ws1.cell(row=start_r1, column=col_idx, value=col_name)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin
    ws1.row_dimensions[start_r1].height = 28
    
    for r_idx, row in df1.iterrows():
        curr_row = start_r1 + 1 + r_idx
        is_ok = (row['Trạng thái'] == 'OK')
        row_fill = fill_ok if is_ok else fill_missing
        row_font = font_ok if is_ok else font_missing
        
        for c_idx, col_name in enumerate(headers1, 1):
            val = row[col_name]
            cell = ws1.cell(row=curr_row, column=c_idx, value=val)
            cell.fill = row_fill
            cell.font = row_font
            cell.border = border_thin
            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.##"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        ws1.row_dimensions[curr_row].height = 20
        
    for col in ws1.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < start_r1:
                continue
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws1.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)
        
    # ===== SHEET 2 =====
    df2 = order_result['detail_table']
    ws2.cell(row=1, column=1, value=f"CHI TIẾT BOM TRÁI & PHẢI - ĐƠN HÀNG: {order_no}").font = font_title
    ws2.cell(row=2, column=1, value=f"Số lượng đơn hàng: {order_qty:g} | Định mức cho từng vế nhân với số lượng đơn").font = font_sub
    
    start_r2 = 4
    headers2 = list(df2.columns)
    for col_idx, col_name in enumerate(headers2, 1):
        cell = ws2.cell(row=start_r2, column=col_idx, value=col_name)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin
    ws2.row_dimensions[start_r2].height = 28
    
    for r_idx, row in df2.iterrows():
        curr_row = start_r2 + 1 + r_idx
        is_ok = (row['Trạng thái'] == 'OK')
        row_fill = fill_ok if is_ok else fill_missing
        row_font = font_ok if is_ok else font_missing
        
        for c_idx, col_name in enumerate(headers2, 1):
            val = row[col_name]
            cell = ws2.cell(row=curr_row, column=c_idx, value=val)
            cell.fill = row_fill
            cell.font = row_font
            cell.border = border_thin
            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.##"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        ws2.row_dimensions[curr_row].height = 20
        
    for col in ws2.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < start_r2:
                continue
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws2.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)
        
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()


def export_all_orders_to_excel(all_results):
    """
    Xuất toàn bộ kết quả các đơn hàng trong danh sách chờ ra một file Excel tổng hợp
    với mỗi đơn hàng gồm 1 sheet tổng hợp vật tư.
    """
    wb = openpyxl.Workbook()
    # Sheet index tóm tắt
    ws_idx = wb.active
    ws_idx.title = "Tong_Ket_Cac_Don_Hang"
    
    font_title = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    
    font_ok = Font(name="Segoe UI", size=10, color="155724")
    fill_ok = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    
    font_missing = Font(name="Segoe UI", size=10, bold=True, color="721C24")
    fill_missing = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    
    border_thin = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    
    ws_idx.cell(row=1, column=1, value="BẢNG TỔNG KẾT ĐỐI SOÁT VẬT TƯ CÁC ĐƠN HÀNG (FIFO)").font = font_title
    headers_idx = ["STT", "Mã đơn hàng", "Sản phẩm", "Số lượng đặt", "Tổng số mã LK", "Số mã ĐỦ", "Số mã THIẾU", "Đánh giá chung"]
    
    for c_i, h in enumerate(headers_idx, 1):
        cell = ws_idx.cell(row=3, column=c_i, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border_thin
    ws_idx.row_dimensions[3].height = 26
    
    for r_i, r in enumerate(all_results, 1):
        curr_r = 3 + r_i
        is_ok = r.get('is_fully_ok', False)
        row_fill = fill_ok if is_ok else fill_missing
        row_font = font_ok if is_ok else font_missing
        
        vals = [
            r['stt'],
            r['order_no'],
            r.get('product_title', ''),
            r['order_qty'],
            r['total_items'],
            r['ok_count'],
            r['missing_count'],
            "ĐỦ VẬT TƯ (OK)" if is_ok else f"THIẾU {r['missing_count']} MÃ"
        ]
        
        for c_i, val in enumerate(vals, 1):
            cell = ws_idx.cell(row=curr_r, column=c_i, value=val)
            cell.fill = row_fill
            cell.font = row_font
            cell.border = border_thin
            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        ws_idx.row_dimensions[curr_r].height = 20
        
    for col in ws_idx.columns:
        max_len = max([len(str(cell.value or '')) for cell in col])
        col_letter = get_column_letter(col[0].column)
        ws_idx.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    # Sheet 2: Chi tiết tất cả các đơn hàng (Trái trước, Phải sau)
    ws_all_det = wb.create_sheet(title="Chi_Tiet_BOM_Tat_Ca_Don")
    ws_all_det.cell(row=1, column=1, value="CHI TIẾT TỔNG HỢP BOM TRÁI & PHẢI TẤT CẢ CÁC ĐƠN HÀNG (FIFO)").font = font_title
    ws_all_det.cell(row=2, column=1, value="Cơ chế FIFO: Ưu tiên đơn hàng trên trước, trừ dần vế Trái trước - vế Phải sau").font = Font(name="Segoe UI", size=10, italic=True, color="4B5563")
    
    headers_all_det = [
        "STT Đơn", "Mã đơn hàng", "Vế", "Mã sản phẩm (Product Code)",
        "ITEM CODE", "DESCRIPTION (tên mô tả)", "Định mức BOM",
        "Số lượng đơn", "Số lượng cần", "Tồn kho khả dụng", "Số lượng cấp",
        "Số lượng thiếu", "Trạng thái", "Ghi chú"
    ]
    
    start_all_det_r = 4
    for c_i, h in enumerate(headers_all_det, 1):
        c = ws_all_det.cell(row=start_all_det_r, column=c_i, value=h)
        c.font = font_header
        c.fill = fill_header
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border_thin
    ws_all_det.row_dimensions[start_all_det_r].height = 28
    
    curr_det_row = start_all_det_r
    for r in all_results:
        df_det = r.get('detail_table', pd.DataFrame())
        if len(df_det) == 0:
            continue
        for _, row_item in df_det.iterrows():
            curr_det_row += 1
            is_ok = (row_item['Trạng thái'] == 'OK')
            row_fill = fill_ok if is_ok else fill_missing
            row_font = font_ok if is_ok else font_missing
            
            row_vals = [
                r['stt'],
                r['order_no'],
                row_item['Vế'],
                row_item['Mã sản phẩm (Product Code)'],
                row_item['ITEM CODE'],
                row_item['DESCRIPTION (tên mô tả)'],
                row_item['Định mức BOM'],
                row_item['Số lượng đơn'],
                row_item['Số lượng cần'],
                row_item['Tồn kho khả dụng'],
                row_item['Số lượng cấp'],
                row_item['Số lượng thiếu'],
                row_item['Trạng thái'],
                row_item['Ghi chú']
            ]
            
            for c_i, val in enumerate(row_vals, 1):
                c = ws_all_det.cell(row=curr_det_row, column=c_i, value=val)
                c.fill = row_fill
                c.font = row_font
                c.border = border_thin
                if isinstance(val, (int, float)):
                    c.alignment = Alignment(horizontal="right", vertical="center")
                    c.number_format = "#,##0.##"
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")
            ws_all_det.row_dimensions[curr_det_row].height = 20
            
    for col in ws_all_det.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < start_all_det_r:
                continue
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws_all_det.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)
        
    # Tạo các sheet con cho từng đơn
    for r in all_results:
        sheet_name = f"DH_{r['order_no']}"[:31]
        ws = wb.create_sheet(title=sheet_name)
        df = r['summary_table']
        
        ws.cell(row=1, column=1, value=f"ĐƠN HÀNG: {r['order_no']} (SL: {r['order_qty']})").font = font_title
        ws.cell(row=2, column=1, value=f"Tổng mã: {r['total_items']} | ĐỦ: {r['ok_count']} | THIẾU: {r['missing_count']}").font = Font(name="Segoe UI", size=10, italic=True)
        
        start_r = 4
        headers = list(df.columns)
        for c_i, h in enumerate(headers, 1):
            c = ws.cell(row=start_r, column=c_i, value=h)
            c.font = font_header
            c.fill = fill_header
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = border_thin
            
        for row_i, row in df.iterrows():
            curr_row = start_r + 1 + row_i
            is_ok = (row['Trạng thái'] == 'OK')
            row_fill = fill_ok if is_ok else fill_missing
            row_font = font_ok if is_ok else font_missing
            
            for c_i, h in enumerate(headers, 1):
                val = row[h]
                c = ws.cell(row=curr_row, column=c_i, value=val)
                c.fill = row_fill
                c.font = row_font
                c.border = border_thin
                if isinstance(val, (int, float)):
                    c.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")
                    
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row < start_r:
                    continue
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)
            
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()


def analyze_orders_readiness(df_bom, df_stock):
    """
    Phân tích đối soát toàn bộ đơn hàng trong BOM đối chiếu với tồn kho hiện có:
    - Tìm ra các đơn hàng có ĐỦ 100% ITEM CODE đồng bộ (tồn kho > 0 cho mọi linh kiện).
    - Tính toán số lượng TỐI ĐA có thể chạy ngay dựa trên linh kiện nghẽn (bottleneck).
    - Liệt kê các đơn chưa đủ mã và chi tiết các mã đang thiếu.
    Tối ưu hóa vector hóa (vectorized groupby) cho tốc độ xử lý tức thì (<0.5s).
    """
    stock_pool = df_stock.groupby('Stock_Code')['Qty'].sum().to_dict()
    orders = get_available_orders_info(df_bom)
    
    # Xác định variant cơ sở cho từng đơn hàng
    base_variants = df_bom.groupby('Order_No')['Variant'].min().to_dict()
    df_bom_base = df_bom[df_bom['Variant'] == df_bom['Order_No'].map(base_variants)].copy()
    
    # Gom nhóm theo Order_No và Component_Code
    grouped = df_bom_base.groupby(['Order_No', 'Component_Code']).agg(
        unit_bom=('Qty', 'sum'),
        desc=('Component_Desc', 'first')
    ).reset_index()
    
    grouped['stock_qty'] = grouped['Component_Code'].map(stock_pool).fillna(0.0)
    grouped['has_stock'] = (grouped['stock_qty'] > 0) & (grouped['unit_bom'] > 0)
    grouped['possible_qty'] = grouped.apply(
        lambda r: (r['stock_qty'] // r['unit_bom']) if r['has_stock'] else 0.0,
        axis=1
    )
    
    ready_orders = []
    not_ready_orders = []
    
    for order_no, oinfo in orders.items():
        sub = grouped[grouped['Order_No'] == order_no]
        total_comps = len(sub)
        comps_in_stock = int(sub['has_stock'].sum())
        
        if comps_in_stock == total_comps and total_comps > 0:
            min_row = sub.sort_values('possible_qty').iloc[0]
            max_runnable = int(min_row['possible_qty'])
            if max_runnable > 0:
                ready_orders.append({
                    'order_no': order_no,
                    'description': oinfo['description'],
                    'total_components': total_comps,
                    'max_runnable_qty': max_runnable,
                    'bottleneck_code': min_row['Component_Code'],
                    'bottleneck_desc': min_row['desc'],
                    'bottleneck_stock': float(min_row['stock_qty']),
                    'bottleneck_unit_bom': float(min_row['unit_bom'])
                })
                continue
                
        # Nếu chưa đủ mã
        missing_rows = sub[~sub['has_stock']]
        not_ready_orders.append({
            'order_no': order_no,
            'description': oinfo['description'],
            'total_components': total_comps,
            'components_in_stock': comps_in_stock,
            'missing_count': len(missing_rows),
            'sample_missing': ', '.join(missing_rows['Component_Code'].iloc[:4].tolist())
        })
        
    ready_orders.sort(key=lambda x: x['max_runnable_qty'], reverse=True)
    not_ready_orders.sort(key=lambda x: x['missing_count'])
    
    total_count = len(orders)
    ready_count = len(ready_orders)
    not_ready_count = len(not_ready_orders)
    ready_percent = round((ready_count / total_count * 100), 1) if total_count > 0 else 0.0
    
    return {
        'ready_orders': ready_orders,
        'not_ready_orders': not_ready_orders,
        'total_orders': total_count,
        'ready_count': ready_count,
        'not_ready_count': not_ready_count,
        'ready_percent': ready_percent
    }


def export_readiness_report_to_excel(readiness_data):
    """
    Xuất báo cáo khả thi sản xuất (đơn hàng đủ mã đồng bộ & số lượng tối đa) ra Excel.
    """
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Don_Du_Ma_San_Sang_Chay"
    ws2 = wb.create_sheet(title="Don_Chua_Du_Ma")
    
    font_title = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
    font_sub = Font(name="Segoe UI", size=10, italic=True, color="4B5563")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    fill_header2 = PatternFill(start_color="475569", end_color="475569", fill_type="solid")
    
    font_ok = Font(name="Segoe UI", size=10, color="155724")
    fill_ok = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    
    border_thin = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    
    # ===== SHEET 1: ĐƠN SẴN SÀNG CHẠY =====
    ws1.cell(row=1, column=1, value="BÁO CÁO CÁC ĐƠN HÀNG ĐỦ MÃ HÀNG ĐỒNG BỘ TRONG TỒN KHO").font = font_title
    ws1.cell(row=2, column=1, value=f"Tổng số đơn đủ mã: {readiness_data['ready_count']} / {readiness_data['total_orders']} đơn ({readiness_data['ready_percent']}%) | Đã tính toán số lượng tối đa có thể sản xuất ngay").font = font_sub
    
    headers1 = [
        "STT", "Mã đơn hàng", "Tên sản phẩm", "Số ITEM CODE trong BOM",
        "SL TỐI ĐA có thể chạy (bộ)", "ITEM CODE điểm nghẽn",
        "Tên linh kiện nghẽn", "Tồn kho linh kiện nghẽn", "ĐM bộ linh kiện nghẽn", "Đánh giá"
    ]
    
    for c_i, h in enumerate(headers1, 1):
        cell = ws1.cell(row=4, column=c_i, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin
    ws1.row_dimensions[4].height = 28
    
    for idx, r in enumerate(readiness_data['ready_orders'], 1):
        curr_row = 4 + idx
        vals = [
            idx,
            r['order_no'],
            r['description'],
            r['total_components'],
            r['max_runnable_qty'],
            r['bottleneck_code'],
            r['bottleneck_desc'],
            r['bottleneck_stock'],
            r['bottleneck_unit_bom'],
            "100% ĐỦ MÃ (Sẵn sàng chạy)"
        ]
        for c_i, val in enumerate(vals, 1):
            cell = ws1.cell(row=curr_row, column=c_i, value=val)
            cell.font = font_ok
            cell.fill = fill_ok
            cell.border = border_thin
            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        ws1.row_dimensions[curr_row].height = 20
        
    for col in ws1.columns:
        max_len = max([len(str(cell.value or '')) for cell in col if cell.row >= 4] or [12])
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)
        
    # ===== SHEET 2: ĐƠN CHƯA ĐỦ MÃ =====
    ws2.cell(row=1, column=1, value="DANH SÁCH ĐƠN HÀNG CHƯA ĐỦ MÃ TRONG TỒN KHO").font = font_title
    ws2.cell(row=2, column=1, value=f"Tổng số đơn chưa đủ mã: {readiness_data['not_ready_count']} đơn").font = font_sub
    
    headers2 = ["STT", "Mã đơn hàng", "Tên sản phẩm", "Tổng ITEM CODE", "Số mã có trong kho", "Số mã còn thiếu", "Mẫu mã linh kiện thiếu"]
    for c_i, h in enumerate(headers2, 1):
        cell = ws2.cell(row=4, column=c_i, value=h)
        cell.font = font_header
        cell.fill = fill_header2
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_thin
    ws2.row_dimensions[4].height = 28
    
    fill_missing = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    font_missing = Font(name="Segoe UI", size=10, color="991B1B")
    
    for idx, r in enumerate(readiness_data['not_ready_orders'], 1):
        curr_row = 4 + idx
        vals = [
            idx,
            r['order_no'],
            r['description'],
            r['total_components'],
            r['components_in_stock'],
            r['missing_count'],
            r['sample_missing']
        ]
        for c_i, val in enumerate(vals, 1):
            cell = ws2.cell(row=curr_row, column=c_i, value=val)
            cell.font = font_missing
            cell.fill = fill_missing
            cell.border = border_thin
            if isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        ws2.row_dimensions[curr_row].height = 20
        
    for col in ws2.columns:
        max_len = max([len(str(cell.value or '')) for cell in col if cell.row >= 4] or [12])
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)
        
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()

